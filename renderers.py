import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BRAND = "0E3B43"


def to_xlsx(report):
    wb = Workbook()
    ws = wb.active
    ws.title = report["title"][:30]
    ws["A1"] = report["title"]
    ws["A1"].font = Font(size=14, bold=True, color=BRAND)
    ws["A2"] = report["subtitle"]
    ws["A3"] = f"Generated {datetime.now():%Y-%m-%d %H:%M}"
    r = 5
    for label, val in report["summary"]:
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2, val)
        r += 1
    r += 1
    for i, col in enumerate(report["columns"], 1):
        c = ws.cell(r, i, col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BRAND)
        c.alignment = Alignment(horizontal="center")
    for row in report["rows"]:
        r += 1
        for i, v in enumerate(row, 1):
            ws.cell(r, i, v)
    for i in range(1, len(report["columns"]) + 1):
        width = max(len(str(report["columns"][i - 1])), *(len(str(row[i - 1])) for row in report["rows"][:200]), 8)
        ws.column_dimensions[get_column_letter(i)].width = min(width + 3, 45)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"


def to_pdf(report):
    buf = io.BytesIO()
    wide = len(report["columns"]) > 6
    page = landscape(A4) if wide else A4
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm,
                            bottomMargin=16 * mm, title=report["title"])
    ss = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=ss["Title"], textColor=colors.HexColor("#" + BRAND), alignment=0, fontSize=18)
    cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=8, leading=10)
    story = [Paragraph(report["title"], h), Paragraph(report["subtitle"], ss["Normal"]),
             Paragraph(f"Generated {datetime.now():%Y-%m-%d %H:%M}", ss["Italic"]), Spacer(1, 6 * mm)]
    if report["summary"]:
        st = Table([[str(k), str(v)] for k, v in report["summary"]], hAlign="LEFT", colWidths=[55 * mm, 40 * mm])
        st.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 9), ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.lightgrey)]))
        story += [st, Spacer(1, 6 * mm)]
    data = [[Paragraph(f"<b>{c}</b>", ParagraphStyle("hd", parent=cell, textColor=colors.white))
             for c in report["columns"]]]
    data += [[Paragraph(str(v), cell) for v in row] for row in report["rows"]] or \
            [[Paragraph("No data for the selected filters.", cell)] + [""] * (len(report["columns"]) - 1)]
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BRAND)),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F4")]),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5D3")),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story.append(t)

    def footer(canvas, d):
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(14 * mm, 8 * mm, "Smart Attendance & Academic Analytics Platform")
        canvas.drawRightString(page[0] - 14 * mm, 8 * mm, f"Page {d.page}")

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue(), "application/pdf", "pdf"
